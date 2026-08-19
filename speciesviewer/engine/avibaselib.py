"""Shared helpers for reading and enriching an AviBase/BirdingBreaks species list."""

import csv
import os
import re
import logging

logger = logging.getLogger("speciesviewer")

DEFAULT_SOURCE_LIST = "avibase.csv"

# What a column in a Birding Breaks list can hold. Anything not marked is
# ignored — that is how the day boxes and spacer columns stay out of the way.
ROLE_KEYS = ("seq", "en", "nl", "sc", "st")
ROLE_LABELS = {
    "seq": "sequence number",
    "en":  "English name",
    "nl":  "Dutch name",
    "sc":  "scientific name",
    "st":  "status or note",
}


def resolve_source_list(triproot, configured=""):
    """Full path of the species list to read for this trip.

    A list picked in the dialog wins; without one, the historical default
    `avibase.csv` in the trip's own data folder. The path is returned whether
    or not it exists, so the caller can name it in an error message.
    """
    configured = (configured or "").strip()
    if configured:
        return os.path.abspath(os.path.expanduser(configured))
    return os.path.join(triproot, DEFAULT_SOURCE_LIST)


def read_avibase(triproot, filename):
    """Read AviBase CSV and return a dict keyed by enname (column 0).

    CSV columns: English/key ; Scientific ; Dutch ; Status
    Section-header rows (no scientific name) get number=-1.
    """
    file_path = os.path.join(triproot, filename)
    number = 1
    avibase_dict = {}

    if not os.path.isfile(file_path):
        # Named in full, because the folder the list is expected in is not the
        # folder the pictures live in — that mix-up is the usual cause.
        raise FileNotFoundError(
            f"No species list found at '{file_path}'. Use Browse… to pick the "
            f"list, or put {DEFAULT_SOURCE_LIST} in the trip data folder "
            f"'{triproot}'."
        )

    with open(file_path, "r", encoding="utf-8") as f:
        for line in f.read().split("\n"):
            tokens = line.split(";")
            try:
                if tokens[0] in ("﻿Column1", "﻿Engels", ""):
                    continue
                if tokens[0].strip() == "":
                    continue
                if len(tokens) > 1 and tokens[1] not in ("Wet", ""):
                    sp = {
                        "enname":       tokens[0],
                        "scname":       tokens[1],
                        "nlname":       tokens[2] if len(tokens) > 2 else "",
                        "status":       tokens[3] if len(tokens) > 3 else "",
                        "english":      "",
                        "number":       number,
                        "gezien":       0,
                        "foto":         0,
                        "knownspecies": False,
                    }
                    number += 1
                    avibase_dict[sp["enname"]] = sp
                elif len(tokens) > 1 and tokens[1] == "Wet":
                    sp = {
                        "enname":       tokens[0],
                        "scname":       tokens[1],
                        "nlname":       tokens[2] if len(tokens) > 2 else "",
                        "status":       tokens[3] if len(tokens) > 3 else "",
                        "english":      "",
                        "number":       "Wet",
                        "gezien":       0,
                        "foto":         0,
                        "knownspecies": False,
                    }
                    avibase_dict[sp["enname"]] = sp
                else:
                    # Section header row — no scientific name
                    avibase_dict[tokens[0]] = {
                        "enname":  tokens[0],
                        "scname":  "",
                        "nlname":  "",
                        "status":  "",
                        "english": "",
                        "number":  -1,
                    }
            except Exception:
                pass

    if not any(entry.get("number") != -1 for entry in avibase_dict.values()):
        # Every line became a section heading, which means the file is not
        # semicolon-separated at all. Saying so beats writing a checklist that
        # holds nothing but grey group rows.
        found = sniff_delimiter(open(file_path, encoding="utf-8-sig",
                                     errors="replace").read().splitlines()[:50])
        named = {"\t": "tabs", ";": "semicolons", ",": "commas",
                 "|": "bars"}.get(found, found)
        raise ValueError(
            f"No species found in '{os.path.basename(file_path)}' with the fixed "
            f"AviBase layout, which expects semicolons between the names. This "
            f"file is separated by {named}. Choose \u201cColumns I marked "
            f"myself\u201d and mark the columns instead."
        )

    return avibase_dict


def add_species_info(avibase_dict, sc, sp):
    """Enrich avibase_dict entries with canonical scientific names and English names.

    sc  — SpeciesClassification instance (has english_dict, dutch_dict, sc_dict)
    sp  — Species instance (has species_dict for known-species check)
    Mutates avibase_dict in place.
    """
    for enname, entry in avibase_dict.items():
        if not entry.get("scname"):
            continue

        original_scname = entry["scname"]

        if enname.lower() in sc.english_dict:
            entry["scname"] = sc.english_dict[enname.lower()]
            if entry["scname"] in sp.species_dict:
                entry["knownspecies"] = True
        else:
            if not entry["nlname"] and entry["scname"].lower() in sc.sc_dict:
                entry["nlname"] = sc.sc_dict[entry["scname"].lower()].get("DutchName", "")
            if entry["nlname"].lower() in sc.dutch_dict:
                entry["scname"] = sc.dutch_dict[entry["nlname"].lower()]
                if entry["scname"] in sp.species_dict:
                    entry["knownspecies"] = True

        # Fill English name from classification when missing
        if not entry.get("english"):
            sc_entry = sc.sc_dict.get(entry["scname"].lower())
            if sc_entry and sc_entry.get("EnglishName"):
                entry["english"] = sc_entry["EnglishName"]

        if (original_scname and entry["scname"] and
                entry["scname"].lower() != original_scname.lower()):
            label = entry.get("nlname") or entry.get("english") or enname
            logger.info(
                f"Scientific name updated for '{label}': "
                f"'{original_scname}' → '{entry['scname']}'"
            )

        if entry["nlname"] == '':
            sc_entry = sc.sc_dict.get(entry["scname"].lower())
            if sc_entry and sc_entry.get("DutchName"):
                entry["nlname"] = sc_entry["DutchName"]



# ── Birding Breaks lists ────────────────────────────────────────────────────
# What arrives from the tour operator is a spreadsheet with title rows above the
# real header and a box per trip day beside each species. Rather than insist on
# one shape, the user marks what each column holds (gui/avibase_columns.py) and
# the reader below follows that map.

def format_roles(roles):
    """Column map as one config line: "0=seq,2=en,3=sc"."""
    return ",".join(f"{col}={role}"
                    for col, role in sorted((roles or {}).items())
                    if role in ROLE_KEYS)


def parse_roles(spec):
    """Column map back from its config line. Unreadable entries are dropped."""
    roles = {}
    for part in (spec or "").split(","):
        if "=" not in part:
            continue
        col, _, role = part.partition("=")
        role = role.strip()
        try:
            col = int(col.strip())
        except ValueError:
            continue
        if role in ROLE_KEYS:
            roles[col] = role
    return roles


def encode_delimiter(delimiter):
    """A delimiter as a config-safe word. An .ini value cannot carry a bare tab
    — it would be written out and read back as nothing at all."""
    if not delimiter:
        return ""
    return "tab" if delimiter == "\t" else delimiter


def decode_delimiter(text):
    """The delimiter back from config. Empty means: work it out from the file."""
    text = (text or "").strip()
    if not text:
        return None
    return "\t" if text.lower() == "tab" else text[0]


def sniff_delimiter(sample_lines):
    """Which delimiter a text list uses. A candidate that appears the same
    number of times on every line beats one that is frequent on a single line."""
    best, best_score = ";", -1
    for cand in (";", ",", "	", "|"):
        counts = [ln.count(cand) for ln in sample_lines if ln.strip()]
        if not counts or max(counts) == 0:
            continue
        score = min(counts) * 10 + sum(counts) / len(counts)
        if score > best_score:
            best, best_score = cand, score
    return best


def _cell_text(value):
    """Spreadsheet cell as the text the user sees — an integer stays "12",
    not "12.0", because that column is usually the sequence number."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _load_spreadsheet_rows(path, limit=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = wb.worksheets
        ws = sheets[0]
        if len(sheets) > 1:
            logger.warning(
                f"'{os.path.basename(path)}' holds {len(sheets)} sheets; reading "
                f"the first one, '{ws.title}'."
            )
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_cell_text(c) for c in row])
            if limit and len(rows) >= limit:
                break
        return rows
    finally:
        wb.close()


def rows_from_text(text, delimiter=None, limit=None):
    """Rows out of delimited text — a file's contents or a paste, same thing."""
    lines = (text or "").splitlines()
    if not delimiter:
        delimiter = sniff_delimiter(lines[:50])
    rows = [[cell.strip() for cell in row]
            for row in csv.reader(lines, delimiter=delimiter)]
    return (rows[:limit] if limit else rows), delimiter


def _load_text_rows(path, delimiter=None, limit=None):
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        return rows_from_text(f.read(), delimiter, limit)


def load_source_rows(path, delimiter=None, limit=None):
    """A species list as rows of plain text, spreadsheet or delimited file alike.

    Returns (rows, delimiter) — the delimiter is None for a spreadsheet, and for
    a text file it is the one sniffed, so the dialog can show what it settled on.
    `limit` reads only the first rows, which is all a preview needs.
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(
            f"No species list found at '{path}'. Use Browse… to pick the list."
        )
    if os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm"):
        return _load_spreadsheet_rows(path, limit), None
    return _load_text_rows(path, delimiter, limit)


def read_birdingbreaks(path, roles, first_data_row=1, delimiter=None):
    """Read a list from disk using the columns the user marked.

    Returns the same dict shape as read_avibase, so everything downstream —
    name enrichment and the Excel writer — is untouched. A list without Dutch
    names is fine: add_species_info fills those from the classification.
    """
    rows, _ = load_source_rows(path, delimiter)
    return parse_rows(rows, roles, first_data_row, os.path.basename(path))


def parse_rows(rows, roles, first_data_row=1, source="the list"):
    """The marked columns turned into species, wherever the rows came from.

    Split from reading a file so that a pasted list can be counted with exactly
    the code that will later read it — a preview that lies is worse than none.
    """
    by_role = {}
    for col, role in (roles or {}).items():
        if role in ROLE_KEYS:
            by_role[role] = int(col)
    if not any(role in by_role for role in ("en", "nl", "sc")):
        raise ValueError(
            "No name column marked in the species list — open Columns… and say "
            "which column holds the English, Dutch or scientific name."
        )

    def cell(row, role):
        idx = by_role.get(role)
        if idx is None or idx >= len(row):
            return ""
        return (row[idx] or "").strip()

    avibase_dict = {}
    number = 1
    species = headings = duplicates = 0

    for row in rows[max(0, first_data_row - 1):]:
        en, nl, sc = cell(row, "en"), cell(row, "nl"), cell(row, "sc")
        status = cell(row, "st")
        if not (en or nl or sc):
            continue
        key = en or sc or nl
        digits = re.sub(r"\D", "", cell(row, "seq"))

        if not sc and not digits:
            # A family or region printed on its own line. Kept, because the
            # checklist uses these to group the species below them.
            avibase_dict[key] = {
                "enname": key, "scname": "", "nlname": "",
                "status": "", "english": "", "number": -1,
            }
            headings += 1
            continue

        if key in avibase_dict:
            duplicates += 1
        seq = int(digits) if digits else number
        avibase_dict[key] = {
            "enname":       key,
            "scname":       sc,
            "nlname":       nl,
            "status":       status,
            "english":      en,
            "number":       seq,
            "gezien":       0,
            "foto":         0,
            "knownspecies": False,
        }
        number = seq + 1
        species += 1

    logger.info(f"Read {species} species and {headings} section headings "
                f"from '{source}', starting at row {first_data_row}.")
    if duplicates:
        logger.warning(f"{duplicates} rows repeat a name already in the list; "
                       f"the last of each is kept.")
    if not species:
        raise ValueError(
            f"No species read from '{source}' — check in Columns… which row the "
            f"species start on and which columns hold the names."
        )
    return avibase_dict
