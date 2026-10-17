"""Ch.6 step 2 — build the new taxonomy tables beside today's list.

Read-only on the master. Writes only into --out:
  taxa.csv      one row per taxon (kingdom -> species): GBIF name + placement,
                a sort number among its siblings + where it came from,
                Dutch/English names as lists (first = shown)
  names.csv     alias table: name -> scientific name it points to, and what kind
  changes.csv   one row per species of today's list: what happens to it
  report.md     what would change against today's list, pictured species first
  preview.txt   the whole tree in the new order
  gbif_cache.json  GBIF species-match answers, so a re-run makes no calls

Sort numbers (compared only among siblings):
  birds         AviList Sequence; orders and families at their lowest member
  hand table    kingdoms, phyla, classes and non-bird orders:
                tools/taxonomy_hand_order.csv (edit it; unknown names come after,
                alphabetically)
  book order    odonata / butterflies: today's family and subfamily numbers
  otherwise     alphabetical

Rules: docs/plans/taxonomy_order_design.md, "Step 2a".

    python tools/svr_taxonomy_build.py --master E:/Dropbox/naturedb \
        --avilist AviList-v2025-11Jun-short.xlsx --out <folder>
"""
import argparse
import csv
import json
import os
import re
import sys
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from difflib import SequenceMatcher
from pathlib import Path

PREFIX_RE = re.compile(r"^(\d+)\s+")
RANKS = ["kingdom", "phylum", "class", "order", "family", "species", "subspecies"]
HIGHER = RANKS[:5]
SAME_STATUS = {"ACCEPTED", "DOUBTFUL"}
KEPT = "kept — GBIF synonym"
KEPT_CHOICE = "kept — your choice"
MERGED_CHOICE = "merged — your choice"
KEPT_REF = "kept — reference list"
OVERRIDES = Path(__file__).with_name("taxonomy_overrides.csv")
HAND_TABLE = Path(__file__).with_name("taxonomy_hand_order.csv")
AFTER_HAND = 1000          # unknown siblings sort after the hand-placed ones


def strip(name):
    return PREFIX_RE.sub("", name or "").strip()


def legacy_number(name):
    """The sort number today's list put in front of a name, or None."""
    m = PREFIX_RE.match(name or "")
    return int(m.group(1)) if m else None


def real_family(name):
    """'zygoptera lestidae' / 'nymphalidae satyrinae' -> the family word in it."""
    words = strip(name).split()
    fam = [w for w in words if w.endswith(("idae", "aceae"))]
    return fam[0] if fam else (words[-1] if words else "")


def similar(a, b):
    """Same word give or take a Latin ending or a spelling (rueppelli/rueppellii,
    crocea/croceus): close overall, or one stem of 4+ letters with a short ending."""
    if SequenceMatcher(None, a, b).ratio() >= 0.8:
        return True
    stem = len(os.path.commonprefix([a, b]))
    return stem >= 4 and len(a) - stem <= 3 and len(b) - stem <= 3   # not damocles/damone


def change_kind(old, new):
    """What GBIF did to one of today's species names."""
    o, n = old.split(), new.split()
    if old == new:
        return "same"
    if len(o) > 2 and o[:2] == n[:2]:
        return "subspecies → species"
    if similar(o[-1] if len(o) == 2 else o[1], n[-1]):
        return "spelling" if o[0] == n[0] else "renamed"
    return "merged"


def read_csv(path, delimiter=";"):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=delimiter))


# ---------------------------------------------------------------- today's list
def load_today(master):
    active = Path(master) / "metadata/active"
    sc = {r["ScientificName"]: r for r in read_csv(active / "speciescollection.csv")}
    pics = {}
    for r in read_csv(active / "species.csv"):
        if r["Level"] == "Level.SPECIES":
            try:
                pics[r["ScientificName"]] = int(r["iPictureCollectionCount"] or 0)
            except ValueError:
                pics[r["ScientificName"]] = 0

    def chain(name):
        out, seen = {}, set()
        while name in sc and name not in seen:
            seen.add(name)
            out[sc[name]["Level"].split(".")[-1].lower()] = name
            name = sc[name]["Up"]
        return out

    species = []
    for name, r in sc.items():
        if r["Level"] != "Level.SPECIES":
            continue
        ch = chain(name)
        if "kingdom" not in ch:          # landscapes and other topics
            continue
        fam_raw = ch.get("family", "")
        species.append({
            "name": name, "nl": r["DutchName"], "en": r["EnglishName"],
            "kingdom": strip(ch.get("kingdom")), "phylum": strip(ch.get("phylum")),
            "class": strip(ch.get("class")), "order": strip(ch.get("order")),
            "family": real_family(fam_raw), "family_raw": strip(fam_raw),
            "book": legacy_number(fam_raw), "pics": pics.get(name, 0),
            "subspecies_of": (r.get("SubspeciesOf") or "").strip().lower(),
        })
    # names of higher taxa, keyed by the name without its number; a fused
    # 'zygoptera lestidae' also names the family 'lestidae'
    higher = {}
    for name, r in sc.items():
        if r["Level"] in ("Level.KINGDOM", "Level.PHYLUM", "Level.CLASS",
                          "Level.ORDER", "Level.FAMILY"):
            higher.setdefault(strip(name), (r["DutchName"], r["EnglishName"]))
            if r["Level"] == "Level.FAMILY" and strip(name).split()[-1] == real_family(name):
                higher.setdefault(real_family(name), (r["DutchName"], r["EnglishName"]))
    return species, higher


def load_overrides(path=OVERRIDES):
    """Your decisions after review (tools/taxonomy_overrides.csv):
      keep      the species stays as it is, whatever GBIF says (merge, rename, spelling)
      merge     the species joins Value (or, without Value, the species GBIF names)
      en / nl   the name to show; the one shown before stays, for searching
      alias-en / alias-nl   an extra name to search by"""
    ov = {"keep": {}, "merge": {}, "names": [], "subspecies": {}, "avilist": {}}
    if not path.exists():
        return ov
    for r in read_csv(path):
        name, act = r["Name"].strip().lower(), r["Action"].strip().lower()
        val, note = r["Value"].strip().lower(), r["Note"].strip()
        if act == "keep":
            ov["keep"].setdefault(name, note)
        elif act == "merge":
            ov["merge"][name] = (val, note)
        elif act in ("subspecies", "avilist"):
            ov[act][name] = (val, note)
        elif act in ("en", "nl", "alias-en", "alias-nl"):
            ov["names"].append((name, act, val, note))
    return ov


def load_hand_table(path=HAND_TABLE):
    return {(r["Rank"], r["Parent"], r["Name"]): (float(r["Sequence"]), r["Source"])
            for r in read_csv(path)}


# ---------------------------------------------------------------- AviList
def load_avilist(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
        hdr = rows[0]
        rows = [dict(zip(hdr, r)) for r in rows[1:]]
    else:
        rows = read_csv(path, delimiter=",")
    species, genus, family_en, by_epithet = {}, {}, {}, defaultdict(list)
    for r in rows:
        rank = (r.get("Taxon_rank") or "").lower()
        name = (r.get("Scientific_name") or "").lower().strip()
        seq = float(r["Sequence"])
        if rank == "species":
            rec = {"name": name, "seq": seq, "en": (r.get("English_name_AviList") or "").lower(),
                   "family": (r.get("Family") or "").lower(), "avibase": r.get("AvibaseID") or ""}
            species[name] = rec
            by_epithet[(name.split()[-1], rec["family"])].append(rec)
        elif rank == "genus":
            genus[name] = seq
        elif rank == "family" and r.get("Family_English_name"):
            family_en[name] = r["Family_English_name"].lower()
    return species, genus, family_en, by_epithet


def load_mdd(path):
    """Mammal Diversity Database (mammaldiversity.org, Zenodo), in the same shape as
    load_avilist. The species order is the ROW order of the file: MDD sorts by
    phylosort (its 27 orders, phylogenetic), then suborder … superfamily, family,
    genus, species — every family one block. phylosort alone ranks only orders."""
    species, genus, by_epithet = {}, {}, defaultdict(list)
    with open(path, encoding="utf-8", newline="") as f:
        for i, r in enumerate(csv.DictReader(f), 1):
            name = r["sciName"].replace("_", " ").lower()
            seq = float(i)
            rec = {"name": name, "seq": seq, "en": (r["mainCommonName"] or "").lower(),
                   "family": (r["family"] or "").lower(), "avibase": ""}
            species[name] = rec
            by_epithet[(name.split()[-1], rec["family"])].append(rec)
            g = name.split()[0]
            genus[g] = min(genus.get(g, seq), seq)
    return species, genus, {}, by_epithet


# ---------------------------------------------------------------- GBIF
def gbif_match(species, cache_path):
    cache = json.loads(cache_path.read_text(encoding="utf-8")) if cache_path.exists() else {}
    todo = [s for s in species if s["name"] not in cache]

    def one(s):
        params = {"name": s["name"], "kingdom": s["kingdom"].capitalize(), "strict": "true"}
        if s["class"] == "aves":
            params["class"] = "Aves"
        q = urllib.parse.urlencode(params)
        for _ in range(3):
            try:
                with urllib.request.urlopen("https://api.gbif.org/v1/species/match?" + q,
                                            timeout=30) as resp:
                    return s["name"], json.load(resp)
            except Exception:
                pass
        return s["name"], None

    if todo:
        print(f"GBIF: {len(todo)} lookups ({len(species) - len(todo)} cached)…", flush=True)
        with ThreadPoolExecutor(8) as ex:
            for i, (name, d) in enumerate(ex.map(one, todo), 1):
                if d is not None:
                    cache[name] = d
                if i % 500 == 0:          # save as we go: an interrupted run resumes
                    cache_path.write_text(json.dumps(cache), encoding="utf-8")
                    print(f"  {i}/{len(todo)}", flush=True)
        cache_path.write_text(json.dumps(cache), encoding="utf-8")
    return cache


def gbif_view(d):
    """GBIF's answer reduced to what we use; None when GBIF has no species for it."""
    if not d or d.get("matchType") == "NONE" or d.get("rank") not in ("SPECIES", "SUBSPECIES"):
        return None
    return {
        "species": (d.get("species") or "").lower(),
        "status": d.get("status"),
        **{k: (d.get(k) or "").lower() for k in ("kingdom", "phylum", "class", "order",
                                                 "family", "genus")},
        "key": d.get("speciesKey") or d.get("usageKey"),
    }


def chain_of(place):
    """[(rank, name)] from kingdom down, skipping empty levels and a name that
    repeats a higher one (GBIF: class Crocodylia, order Crocodylia)."""
    out, seen = [], set()
    for rank in HIGHER:
        name = place.get(rank, "")
        if name and name not in seen:
            out.append((rank, name))
            seen.add(name)
    return out



# ---------------------------------------------------------------- authorities
# Birds come from AviList and mammals from MDD: names, placement and order. GBIF
# is never asked about them (decided 16 Oct). Both lists are CC BY 4.0.
def _norm_en(s):
    return (s or "").lower().replace("-", " ").replace("grey", "gray").strip()


def _new_ref(label, klass):
    return {"label": label, "klass": klass, "species": {}, "subspecies": {}, "genus": {},
            "by_en": defaultdict(set), "synonyms": defaultdict(set),
            "by_ep_fam": defaultdict(set), "family_en": {}}


def load_avilist_ref(path):
    """AviList EXTENDED xlsx -> a reference: species, subspecies, genera, and what
    today's names can be matched by (English names in three columns, protonym)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    h = next(it)
    ix = {k: i for i, k in enumerate(h)}

    def col(r, k):
        return str(r[ix[k]]) if k in ix and r[ix[k]] is not None else ""
    ref = _new_ref("AviList", "aves")
    for r in it:
        rank, name = col(r, "Taxon_rank").lower(), col(r, "Scientific_name").lower().strip()
        seq = float(r[ix["Sequence"]])
        order, family = col(r, "Order").lower(), col(r, "Family").lower()
        if rank == "species":
            ref["species"][name] = {"seq": seq, "order": order, "family": family,
                                    "en": col(r, "English_name_AviList").lower(),
                                    "id": col(r, "AvibaseID")}
            for c in ("English_name_AviList", "English_name_Clements_v2024",
                      "English_name_BirdLife_v9"):
                if col(r, c):
                    ref["by_en"][_norm_en(col(r, c))].add(name)
            proto = " ".join(col(r, "Protonym").lower().split()[:2])
            if proto and proto != name:
                ref["synonyms"][proto].add(name)
            ref["by_ep_fam"][(name.split()[-1], family)].add(name)
        elif rank == "subspecies":
            ref["subspecies"][name] = {"seq": seq, "species": " ".join(name.split()[:2])}
        elif rank == "genus":
            ref["genus"][name] = {"seq": seq, "order": order, "family": family}
        elif rank == "family" and col(r, "Family_English_name"):
            ref["family_en"][family] = col(r, "Family_English_name").lower()
    return ref


def load_mdd_ref(path, syn_path=None):
    """MDD species CSV (+ the Species_Syn CSV for old names) -> a reference. The
    species order is the file's row order (27 orders phylogenetic, families grouped)."""
    ref = _new_ref("MDD", "mammalia")
    with open(path, encoding="utf-8", newline="") as f:
        for i, r in enumerate(csv.DictReader(f), 1):
            name = r["sciName"].replace("_", " ").lower()
            order, family = r["order"].lower(), r["family"].lower()
            ref["species"][name] = {"seq": float(i), "order": order, "family": family,
                                    "en": (r["mainCommonName"] or "").lower(), "id": r["id"]}
            for en in [r["mainCommonName"]] + (r["otherCommonNames"] or "").split("|"):
                if en and en != "NA":
                    ref["by_en"][_norm_en(en)].add(name)
            ref["by_ep_fam"][(name.split()[-1], family)].add(name)
            g = name.split()[0]
            if g not in ref["genus"]:
                ref["genus"][g] = {"seq": float(i), "order": order, "family": family}
    if syn_path:
        with open(syn_path, encoding="utf-8", newline="") as f:
            for r in csv.DictReader(f):
                old = (r["MDD_normalized_original_combination"] or "").lower().strip()
                acc = (r["MDD_species"] or "").replace("_", " ").lower().strip()
                if old and acc in ref["species"] and len(old.split()) == 2 and old != acc:
                    ref["synonyms"][old].add(acc)
    return ref


def resolve_authority(species_today, ref, ov):
    """Today's names in one class, against its reference list. Returns rows (every
    species of the list + today's kept ones), subs (subspecies rows), names, changes."""
    L, sp, klass = ref["label"], ref["species"], ref["klass"]
    rows, subs, names, changes = {}, {}, [], []

    def place_of(order, family):
        return {"kingdom": "animalia", "phylum": "chordata", "class": klass,
                "order": order, "family": family}

    for name, a in sp.items():
        rows[name] = {"name": name, "chain": chain_of(place_of(a["order"], a["family"])),
                      "bird": True, "seq": a["seq"], "seq_src": L, "book": None,
                      "book_family": a["family"], "nl": [], "en": [a["en"]] if a["en"] else [],
                      "avibase": a["id"] if L == "AviList" else "", "gbif_key": "",
                      "pics": 0, "merged_from": [], "note": ""}
    exact = {s["name"] for s in species_today if s["name"] in sp}
    sub_by_gen_ep = {(n.split()[0], n.split()[-1]): v["species"]
                     for n, v in ref["subspecies"].items()}

    def unique(cands):
        c = set(cands) - exact
        return next(iter(c)) if len(c) == 1 else None

    quiet = ("same", "same English name", "subspecies — your choice",
             "alias — your choice", MERGED_CHOICE)
    for s in sorted(species_today, key=lambda s: s["name"] not in exact):
        n, w = s["name"], s["name"].split()
        target, sub, kind, decision = None, None, "", ""
        by_en = unique(ref["by_en"].get(_norm_en(s["en"]), ())) if s["en"] else None
        if s.get("subspecies_of") and s["subspecies_of"] in sp:
            sub, target, kind = n, s["subspecies_of"], "subspecies — already converted"
        elif n in ov["merge"]:
            (target, decision), kind = ov["merge"][n], MERGED_CHOICE
        elif n in ov["subspecies"]:
            sub, decision = ov["subspecies"][n]
            target, kind = " ".join(sub.split()[:2]), "subspecies — your choice"
        elif n in ov["avilist"]:
            (target, decision), kind = ov["avilist"][n], "alias — your choice"
        elif n in sp:
            target, kind = n, "same"
        elif n in ov["keep"]:
            kind, decision = KEPT_CHOICE, ov["keep"][n]
        elif len(w) == 3 and " ".join(w[:2]) in sp:
            target, kind = " ".join(w[:2]), "subspecies → species"
        elif (w[0], w[-1]) in sub_by_gen_ep:
            target, kind = sub_by_gen_ep[(w[0], w[-1])], f"a subspecies in {L} → species"
        elif by_en:
            target, kind = by_en, "same English name"
        elif unique(ref["synonyms"].get(n, ())):
            target, kind = unique(ref["synonyms"][n]), f"synonym in {L}"
        else:
            fam = ref["genus"].get(w[0], {}).get("family") or s["family"]
            target = unique(ref["by_ep_fam"].get((w[-1], fam), ()))
            kind = "epithet in family — REVIEW" if target else f"not in {L} — REVIEW"
        if target and target not in sp:
            kind, target = f"not in {L} — REVIEW (no {target})", None

        if target is None:                       # kept as its own species
            g = ref["genus"].get(w[0], {})
            seq = next((v["seq"] for sn, v in ref["subspecies"].items()
                        if sn.split()[0] == w[0] and sn.split()[-1] == w[-1]), None)
            if seq is None and g:
                seq = g["seq"] + 0.5
            order, family = g.get("order") or s["order"], g.get("family") or s["family"]
            rows[n] = {"name": n, "chain": chain_of(place_of(order, family)), "bird": True,
                       "seq": seq, "seq_src": f"own species, placed by {L}" if seq else "end of family",
                       "book": None, "book_family": family,
                       "nl": [s["nl"]] if s["nl"] else [], "en": [s["en"]] if s["en"] else [],
                       "avibase": "", "gbif_key": "", "pics": s["pics"], "merged_from": [],
                       "note": "; ".join(x for x in (kind, decision) if x)}
            changes.append({**s, "new": n, "kind": kind, "new_order": order,
                            "new_family": family, "new_class": klass,
                            "seq_src": rows[n]["seq_src"], "gbif": "", "decision": decision,
                            "reference": L})
            continue

        row = rows[target]
        if sub:                                   # its own subspecies row under the species
            v = ref["subspecies"].get(sub, {})
            subs[sub] = {"name": sub, "parent": target, "seq": v.get("seq", row["seq"] + 0.0001),
                         "seq_src": L if v else "after its species",
                         "nl": [s["nl"]] if s["nl"] else [], "en": [s["en"]] if s["en"] else [],
                         "pics": s["pics"], "note": decision, "avibase": "", "gbif_key": ""}
            new = sub
        else:
            new = target
            row["pics"] += s["pics"]
            if s["nl"] and s["nl"] not in row["nl"]:
                row["nl"].append(s["nl"])
                if len(row["nl"]) > 1:
                    names.append([s["nl"], "nl", target, kind, "today's list"])
            if s["en"] and s["en"] not in row["en"]:
                row["en"].append(s["en"])
                names.append([s["en"], "en", target, f"today's English name ({kind})",
                              "today's list"])
        if n != new and n not in sp:       # a name the list itself uses (the extinct
            k = kind + (" — REVIEW" if s["pics"] and "REVIEW" not in kind and kind not in quiet
                        else "")             # aurochs) keeps its meaning: no alias
            names.append([n, "sci", new, k, L])
        changes.append({**s, "new": new, "kind": kind, "new_order": sp[target]["order"],
                        "new_family": sp[target]["family"], "new_class": klass, "seq_src": L,
                        "gbif": "", "decision": decision, "reference": L})
    return rows, subs, names, changes


# ---------------------------------------------------------------- build
def build(species, higher, avi, gbif, hand, ov, mdd=None, refcheck=False, auth=None):
    """refcheck: follow a GBIF genus move only when the group's reference list
    (AviList for birds, MDD for mammals) uses the new name too."""
    av_family_en = avi[2]
    refs = {"aves": ("AviList", avi)}
    if mdd:
        refs["mammalia"] = ("MDD", mdd)

    def ref_keeps(old, new, klass):
        """The reference list still uses today's name and not GBIF's: GBIF is behind."""
        if not refcheck or klass not in refs:
            return ""
        label, (sp, *_rest) = refs[klass]
        return f"{label} still uses this name" if old in sp and new not in sp else ""
    rows = {}          # new species name -> row
    names = []         # alias rows
    changes = []       # per old species: what happens to it
    subs = {}          # subspecies rows (birds / mammals only)
    if auth:           # birds and mammals from their own list; GBIF only for the rest
        for klass, ref in auth.items():
            r_, s_, n_, c_ = resolve_authority([s for s in species if s["class"] == klass],
                                               ref, ov)
            rows.update(r_)
            subs.update(s_)
            names.extend(n_)
            changes.extend(c_)
            av_family_en = {**ref["family_en"], **av_family_en}
        species = [s for s in species if s["class"] not in auth]

    def new_name(s):
        g = gbif_view(gbif.get(s["name"]))
        if s["name"] in ov["merge"]:
            return ov["merge"][s["name"]][0] or (g["species"] if g else s["name"])
        if s["name"] in ov["keep"]:
            return s["name"]
        if not g or not g["species"] or change_kind(s["name"], g["species"]) == "merged":
            return s["name"]    # GBIF merges are never followed: the species stays
        if ref_keeps(s["name"], g["species"], g["class"] or s["class"]):
            return s["name"]
        return g["species"]

    # species that exist in the new table under their own name: an epithet guess
    # must never land on one of them (Oxyura ferruginea is not Tadorna ferruginea)
    taken = {new_name(s) for s in species}

    # a species GBIF leaves alone owns its row; one merging into it comes second
    for s in sorted(species, key=lambda s: new_name(s) != s["name"]):
        g = gbif_view(gbif.get(s["name"]))
        today = {k: s[k] for k in HIGHER}
        if g and g["species"]:
            new = g["species"]
            kind = change_kind(s["name"], new)
            if kind == "merged":
                # GBIF calls it a synonym of another species. Never followed (too
                # often wrong: Ochlodes sylvanus -> venata); it stays as it is.
                kind, gbif_target, new = KEPT, new, s["name"]
            place = {k: g[k] for k in HIGHER}
            if not place["order"] and not place["family"]:
                place["order"], place["family"] = today["order"], today["family"]
            elif not place["family"]:
                place["family"] = today["family"]   # GBIF leaves it empty: today's stays
            genus = g["genus"]
        else:
            new, kind, place, genus = s["name"], "not in GBIF", today, s["name"].split()[0]
        if kind != KEPT:
            gbif_target = ""
        reference = ref_keeps(s["name"], new, place["class"]) if kind == "renamed" else ""
        if reference:
            kind, gbif_target, new = KEPT_REF, new, s["name"]
        decision = ""
        if s["name"] in ov["merge"]:                # your choice after review
            target, decision = ov["merge"][s["name"]]
            new, kind, gbif_target = target or gbif_target or new, MERGED_CHOICE, ""
        elif s["name"] in ov["keep"]:
            decision = ov["keep"][s["name"]]
            if gbif_target or new != s["name"]:
                gbif_target, new, kind = gbif_target or new, s["name"], KEPT_CHOICE
        ref = refs.get(place["class"])
        bird = ref is not None      # has a species-level reference order (birds, mammals)

        # sort number from the reference list; other species are numbered per family below
        seq, src, av = None, "", None
        if bird:
            label, (av_species, av_genus, _fe, av_epithet) = ref
            if new in av_species:
                seq, src = av_species[new]["seq"], label
            elif s["name"] in av_species:
                seq, src = av_species[s["name"]]["seq"], f"{label} (today's name)"
            else:
                cands = av_epithet.get((new.split()[-1], place["family"]), []) or \
                        av_epithet.get((s["name"].split()[-1], place["family"]), [])
                cands = [c for c in cands if c["name"] not in taken]
                if len(cands) == 1:
                    seq, src = cands[0]["seq"], f"{label} epithet in family — REVIEW"
                elif genus in av_genus:
                    seq, src = av_genus[genus] + 0.5, "after its genus"
                else:
                    src = "none — end of family"
            av = av_species.get(new) or av_species.get(s["name"])

        changes.append({**s, "new": new, "kind": kind, "new_order": place["order"],
                        "new_family": place["family"], "new_class": place["class"],
                        "seq_src": src, "gbif": gbif_target, "decision": decision,
                        "reference": reference})

        if new in rows:                     # a second of today's species lands here
            row = rows[new]
            row["merged_from"].append(s["name"])
            k = kind + (" — REVIEW" if s["pics"] else "")
            names.append([s["name"], "sci", new, k, "GBIF match"])
            for lang in ("nl", "en"):
                if s[lang] and s[lang] not in row[lang]:
                    row[lang].append(s[lang])
                    names.append([s[lang], lang, new, k, "today's list"])
            row["pics"] += s["pics"]
            if s["book"] is not None and (row["book"] is None or s["book"] < row["book"]):
                row["book"] = s["book"]
            continue

        en = [s["en"]] if s["en"] else []
        if av and av["en"] and av["en"] not in en:
            en.append(av["en"])
            names.append([av["en"], "en", new, "other English name", ref[0]])
        rows[new] = {"name": new, "chain": chain_of(place), "bird": bird,
                     "seq": seq, "seq_src": src, "book": None if bird else s["book"],
                     "book_family": s["family"],
                     "nl": [s["nl"]] if s["nl"] else [], "en": en,
                     "avibase": av["avibase"] if av else "", "gbif_key": g["key"] if g else "",
                     "pics": s["pics"], "merged_from": [],
                     "note": "; ".join(x for x in (
                         f"GBIF: {gbif_target}" if gbif_target else "",
                         reference,
                         f"your choice: {decision}" if decision else "") if x)}
        if gbif_target:
            rows[new]["gbif_key"] = ""      # GBIF's key is the other species'
        if new != s["name"]:
            names.append([s["name"], "sci", new, kind, "GBIF match"])

    # names you chose: shown first; the one shown before stays, for searching —
    # unless it is (now) another species' name: then it would find the wrong one
    claimed = defaultdict(set)              # (lang, name) -> species it names by choice
    dropped, pending = [], []
    for name, act, val, note in ov["names"]:
        if act in ("en", "nl") and val:
            claimed[(act, val)].add(name)
    taxon = lambda n: rows.get(n) or subs.get(n)
    for name, act, val, note in ov["names"]:
        if not taxon(name) or not val:
            continue
        lang = act[-2:]
        lst = taxon(name)[lang]
        if act.startswith("alias"):
            if val not in lst:
                lst.append(val)
                names.append([val, lang, name, "other name — your choice", note or "your list"])
            continue
        before = lst[0] if lst else ""
        if val in lst:
            lst.remove(val)
        lst.insert(0, val)
        if before and before != val:
            pending.append((name, lang, before))
    # decided after ALL names are applied, so the file's row order doesn't matter
    for name, lang, before in pending:
        others = (claimed.get((lang, before), set()) |
                  {n for n, r in list(rows.items()) + list(subs.items())
                   if r[lang][:1] == [before]}) - {name}
        if others:
            taxon(name)[lang].remove(before)
            dropped.append((name, lang, before, sorted(others)))
        else:
            names.append([before, lang, name, "name shown before", "today's list"])

    # a subspecies without a name of its own (or with just its species' name):
    # "<species name> (ssp <epithet>)", like "kalkoengier (ssp ruficollis)"
    for name, sb in subs.items():
        species_row = rows.get(sb["parent"], {})
        for lang in ("en", "nl"):
            species_names = species_row.get(lang, [])
            # a name the species also has finds the species: not the subspecies'
            sb[lang] = [n for n in sb[lang] if n not in species_names]
            if not sb[lang] and species_names:
                sb[lang] = [f"{species_names[0]} (ssp {name.split()[-1]})"]

    for r in rows.values():
        r["parent"] = r["chain"][-1][1] if r["chain"] else ""

    # ---- species order inside each family
    by_family = defaultdict(list)
    for r in rows.values():
        by_family[r["parent"]].append(r)
    for fam, members in by_family.items():
        birds = [r for r in members if r["bird"]]
        if birds:
            known = [r["seq"] for r in birds if r["seq"] is not None]
            for r in birds:
                if r["seq"] is None:
                    r["seq"] = (max(known) + 0.9) if known else 999999.0
            tied = defaultdict(list)       # several after the same genus: alphabetical
            for r in birds:
                tied[r["seq"]].append(r)
            for group in tied.values():
                if len(group) > 1:
                    for i, r in enumerate(sorted(group, key=lambda r: r["name"])):
                        r["seq"] += i * 0.001
        others = sorted((r for r in members if not r["bird"]),
                        key=lambda r: (r["book"] if r["book"] is not None else 0, r["name"]))
        for i, r in enumerate(others, 1):
            r["seq"] = float(i)
            r["seq_src"] = "today's book order, then alphabetical" if r["book"] is not None \
                else "alphabetical"

    # ---- higher taxa
    taxa_hi = {}       # name -> {rank, parent, members, books}
    for r in rows.values():
        parent = ""
        for rank, name in r["chain"]:
            t = taxa_hi.setdefault(name, {"rank": rank, "parent": parent, "seqs": [],
                                          "books": [], "bird": False})
            t["seqs"].append(r["seq"])
            t["bird"] |= r["bird"]
            # a family takes a book number only from species that were in it
            # before — not from one GBIF moved in from another family
            if r["book"] is not None and (rank != "family" or name == r["book_family"]):
                t["books"].append(r["book"])
            parent = name
    siblings = defaultdict(list)
    for name, t in taxa_hi.items():
        siblings[t["parent"]].append(name)
    for parent, kids in siblings.items():
        later = []
        for name in kids:
            t = taxa_hi[name]
            if t["bird"] and t["rank"] in ("order", "family"):
                t["seq"], t["src"] = min(t["seqs"]), "lowest member"
            elif (t["rank"], parent, name) in hand:
                t["seq"], t["src"] = hand[(t["rank"], parent, name)]
            elif t["rank"] == "family" and t["books"]:
                t["seq"], t["src"] = float(min(t["books"])), "today's book order"
            else:
                later.append(name)
        for i, name in enumerate(sorted(later), 1):
            taxa_hi[name]["seq"] = AFTER_HAND + i
            taxa_hi[name]["src"] = "alphabetical"

    taxa = []
    for name, t in taxa_hi.items():
        nl, en = higher.get(name, ("", ""))
        if not en and t["rank"] == "family":
            en = av_family_en.get(name, "")
        taxa.append({"rank": t["rank"], "name": name, "parent": t["parent"],
                     "seq": t["seq"], "seq_src": t["src"],
                     "nl": [nl] if nl else [], "en": [en] if en else [],
                     "avibase": "", "gbif_key": "", "pics": "", "note": ""})
    for r in rows.values():
        taxa.append({"rank": "species", **r})
    for r in subs.values():
        taxa.append({"rank": "subspecies", **r})
    for name, lang, before, others in dropped:
        print(f"  name dropped: '{before}' ({lang}) no longer finds {name} — it names {', '.join(others)}")
    return taxa, rows, names, changes


# ---------------------------------------------------------------- output
def write_outputs(out, taxa, rows, names, changes, av_species):
    with open(out / "taxa.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Rank", "ScientificName", "Parent", "Sequence", "SequenceSource",
                    "DutchNames", "EnglishNames", "AvibaseID", "GBIFKey", "Pictures", "Note"])
        for t in sorted(taxa, key=lambda t: (RANKS.index(t["rank"]), t["parent"], t["seq"])):
            w.writerow([t["rank"], t["name"], t["parent"], round(t["seq"], 4), t["seq_src"],
                        t["nl"], t["en"], t["avibase"], t["gbif_key"], t["pics"], t["note"]])
    with open(out / "names.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Name", "Language", "PointsTo", "Kind", "Source"])
        w.writerows(sorted(names, key=lambda n: (n[3], n[0])))

    with open(out / "changes.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Today", "Class", "OrderToday", "FamilyToday", "Kind", "New", "GBIFSynonymOf",
                    "OrderNew", "FamilyNew", "DutchName", "EnglishName", "Pictures", "SequenceSource",
                    "Decision", "Reference"])
        for c in sorted(changes, key=lambda c: (c["class"], c["name"])):
            w.writerow([c["name"], c["class"], c["order"], c["family"], c["kind"], c["new"], c["gbif"],
                        c["new_order"], c["new_family"], c["nl"], c["en"], c["pics"], c["seq_src"],
                        c["decision"], c["reference"]])

    # preview: the tree in the new order
    by_parent = defaultdict(list)
    for t in taxa:
        by_parent[t["parent"]].append(t)
    lines = []

    def walk(parent, depth):
        for t in sorted(by_parent.get(parent, []), key=lambda t: t["seq"]):
            label = t["en"][0] if t["en"] else (t["nl"][0] if t["nl"] else "")
            pics = f"  [{t['pics']}]" if t["rank"] in ("species", "subspecies") and t["pics"] else ""
            lines.append(f"{'  ' * depth}{t['name']}  — {label}{pics}")
            if t["rank"] != "subspecies":
                walk(t["name"], depth + 1)
    walk("", 0)
    (out / "preview.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")

    # report
    pictured = [c for c in changes if c["pics"]]
    kinds_all, kinds_pic = Counter(c["kind"] for c in changes), Counter(c["kind"] for c in pictured)
    fam_moves = [c for c in pictured if c["new_family"] != c["family"]]
    spelled = [c for c in fam_moves if similar(c["family"], c["new_family"])
               and (c["new_order"] == c["order"] or similar(c["order"], c["new_order"]))]
    fam_moves = [c for c in fam_moves if c not in spelled]
    auto = [c for c in fam_moves if c["new_order"] == c["order"]]
    review = [c for c in fam_moves if c["new_order"] != c["order"]]
    landing = defaultdict(list)
    for c in pictured:
        landing[c["new"]].append(c)
    two_into_one = [cs for cs in landing.values() if len(cs) > 1]
    missing = sorted(set(av_species) - {c["new"] for c in changes} - {c["name"] for c in changes})

    L = ["# Step 2 — new taxonomy tables against today's list", "",
         f"Today's list: {len(changes)} species, {len(pictured)} with pictures. "
         f"New table: {len(rows)} species.", "",
         "| class | species | pictured | renamed | kept: GBIF synonym | not in GBIF |",
         "|---|---:|---:|---:|---:|---:|"]
    per_class = defaultdict(list)
    for c in changes:
        per_class[c["class"]].append(c)
    for k, cs in sorted(per_class.items(), key=lambda kv: -len(kv[1])):
        kc = Counter(c["kind"] for c in cs)
        L.append(f"| {k} | {len(cs)} | {sum(1 for c in cs if c['pics'])} | "
                 f"{kc['renamed'] + kc['spelling']} | {kc[KEPT] + kc[KEPT_CHOICE]} | {kc['not in GBIF']} |")
    L += ["", "| what happens | all | with pictures |", "|---|---:|---:|"]
    for k, n in kinds_all.most_common():
        L.append(f"| {k} | {n} | {kinds_pic[k]} |")

    def table(title, items, cols):
        nonlocal L
        L += ["", f"## {title} ({len(items)})", "", "| " + " | ".join(h for h, _ in cols) + " |",
              "|" + "---|" * len(cols)]
        for c in items:
            L.append("| " + " | ".join(str(fn(c)) for _, fn in cols) + " |")

    name_of = lambda c: c["en"] or c["nl"]
    table("Two of your pictured entries become one (a subspecies folded in) — ask the user", two_into_one,
          [("GBIF species", lambda cs: cs[0]["new"]),
           ("today", lambda cs: ", ".join(f"{c['name']} ({name_of(c)}, {c['pics']})" for c in cs))])
    table("Pictured species GBIF calls a synonym — kept as their own species; ask the user",
          sorted([c for c in pictured if c["kind"] == KEPT], key=lambda c: -c["pics"]),
          [("species", lambda c: c["name"]), ("name", name_of), ("GBIF: synonym of", lambda c: c["gbif"]),
           ("pictures", lambda c: c["pics"])])
    table("Pictured species with today's subspecies folded into the species",
          [c for c in pictured if c["kind"] == "subspecies → species"],
          [("today", lambda c: c["name"]), ("GBIF", lambda c: c["new"]), ("pictures", lambda c: c["pics"])])
    table("Pictured species, name spelled differently (Latin ending, spelling)",
          [c for c in pictured if c["kind"] == "spelling"],
          [("today", lambda c: c["name"]), ("GBIF", lambda c: c["new"]), ("pictures", lambda c: c["pics"])])
    table("Pictured species renamed (genus move) — ask the user",
          sorted([c for c in pictured if c["kind"] == "renamed"], key=lambda c: -c["pics"]),
          [("today", lambda c: c["name"]), ("GBIF", lambda c: c["new"]), ("pictures", lambda c: c["pics"])])
    table("Pictured species moving to another ORDER — review", review,
          [("species", lambda c: c["new"]), ("order", lambda c: f"{c['order']} → {c['new_order']}"),
           ("family", lambda c: f"{c['family']} → {c['new_family']}"), ("pictures", lambda c: c["pics"])])
    sp = Counter((f"{c['order']} / {c['family']}", f"{c['new_order']} / {c['new_family']}") for c in spelled)
    L += ["", f"## Order/family only spelled differently by GBIF ({len(spelled)} pictured species)", "",
          "| today → GBIF | species |", "|---|---:|"]
    L += [f"| {a} → {b} | {n} |" for (a, b), n in sp.most_common()]
    moves = Counter((c["family"], c["new_family"]) for c in auto)
    L += ["", f"## Family changes inside the same order — automatic ({len(auto)} pictured species)", "",
          "| family today → new | species |", "|---|---:|"]
    L += [f"| {a} → {b} | {n} |" for (a, b), n in moves.most_common()]
    table("Pictured species not in GBIF — keep today's placement",
          [c for c in pictured if c["kind"] == "not in GBIF"],
          [("species", lambda c: c["name"]), ("class", lambda c: c["class"]),
           ("family", lambda c: c["family"]), ("pictures", lambda c: c["pics"])])
    table("Pictured birds sorted by a guess — review",
          [c for c in pictured if "REVIEW" in c["seq_src"] or "after" in c["seq_src"]
           or "none" in c["seq_src"]],
          [("species", lambda c: c["new"]), ("how", lambda c: c["seq_src"])])
    L += ["", f"## AviList species not in today's list: {len(missing)}", "",
          "(Neither under today's name nor under GBIF's — what a full bird list would add.)", ""]
    (out / "report.md").write_text("\n".join(L) + "\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--master", required=True)
    ap.add_argument("--avilist", required=True, help="AviList short .xlsx (or a .csv of it)")
    ap.add_argument("--out", required=True)
    ap.add_argument("--mdd", help="Mammal Diversity Database CSV (MDD_v2.x_*species.csv)")
    ap.add_argument("--mdd-syn", help="MDD Species_Syn CSV: old mammal names -> MDD species")
    ap.add_argument("--authority", action="store_true",
                    help="birds from AviList (EXTENDED xlsx as --avilist), mammals from MDD (--mdd);"
                         " GBIF only for the other groups")
    ap.add_argument("--reference-check", action="store_true",
                    help="follow a GBIF genus move only when AviList / MDD agree")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)

    species, higher = load_today(a.master)
    avi = load_avilist(a.avilist)
    gbif = gbif_match([s for s in species if not (a.authority and s["class"] in
                                                  ("aves", "mammalia"))], out / "gbif_cache.json")
    mdd = load_mdd(a.mdd) if a.mdd else None
    auth = None
    if a.authority:
        auth = {"aves": load_avilist_ref(a.avilist)}
        if a.mdd:
            auth["mammalia"] = load_mdd_ref(a.mdd, a.mdd_syn)
    taxa, rows, names, changes = build(species, higher, avi, gbif, load_hand_table(),
                                       load_overrides(), mdd, a.reference_check, auth)
    write_outputs(out, taxa, rows, names, changes, avi[0])
    print(f"{len(species)} species today -> {len(rows)} species, {len(taxa)} taxa, "
          f"{len(names)} alias rows. Written to {out}")


if __name__ == "__main__":
    sys.exit(main())
