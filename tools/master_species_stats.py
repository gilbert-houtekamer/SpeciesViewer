"""Master picture statistics per species → Excel (READ-ONLY on the master).

Answers "where do I take too many pictures?": pictures per species, per
species-day (one outing's burst), per camera and per year, with average /
median / percentiles. Reads metadata/active/data_pictures.csv; writes only
the .xlsx.

A *picture* is one UniquePictureName (a shot: RAW + its JPG/sidecars). *Files*
is everything listed under it (RAW, JPG, DNG, XMP …).

    .venv/Scripts/python.exe tools/master_species_stats.py [master_root] [out.xlsx]
"""
import ast
import csv
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BUCKETS = [(1, 1), (2, 5), (6, 10), (11, 25), (26, 50), (51, 100),
           (101, 250), (251, 500), (501, 10**9)]


def _master_root() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "speciesviewer"))
    from configparser import ConfigParser
    from utils import ini_io  # noqa: E402  (project's own .ini reader)
    cfg = ConfigParser(interpolation=None)
    from utils.runtime_files import svconfig_path
    ini_io.read_ini(cfg, svconfig_path())
    return Path(cfg.get("master", "masterbiotadatabase"))


def _stats(values) -> dict:
    v = sorted(values)
    if not v:
        return {}
    q = statistics.quantiles(v, n=20) if len(v) > 1 else [v[0]] * 19
    return {"n": len(v), "total": sum(v), "mean": round(statistics.fmean(v), 1),
            "median": statistics.median(v), "p75": q[14], "p90": q[17],
            "p95": q[18], "max": v[-1]}


def _taxa(target_path: str) -> tuple:
    """biota/<kingdom>/<phylum>/<class>/<order>/<family>/<species>/ → (class, order, family).
    The Kingdom..Family CSV columns are shifted by one, so the path is the truth."""
    parts = [p for p in target_path.replace("\\", "/").split("/") if p]
    return (parts[3], parts[4], parts[5]) if len(parts) >= 7 else ("", "", "")


def _files(cell: str) -> int:
    try:
        return len(ast.literal_eval(cell)) if cell else 0
    except (ValueError, SyntaxError):
        return 0


def main():
    root = _master_root()
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path.home() / "Documents" / "master_species_stats.xlsx"
    src = root / "metadata" / "active" / "data_pictures.csv"
    with open(src, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh, delimiter=";"))

    sp = defaultdict(lambda: {"pics": 0, "files": 0, "days": Counter(), "cams": Counter(),
                              "years": Counter(), "eng": "", "nl": "", "taxa": ("", "", ""),
                              "first": "", "last": ""})
    cam_species = defaultdict(Counter)       # camera → species → pictures
    cam_day = defaultdict(Counter)           # camera → (species, day) → pictures
    year_cam = defaultdict(Counter)          # year → camera → pictures
    burst = Counter()                        # (species, day, camera) → pictures
    no_species = 0
    for r in rows:
        sci = r["WetSoortNaam"].strip().lower()
        if not sci:
            no_species += 1
            continue
        cam = r["CameraModel"].strip() or "(unknown)"
        dt = r["DateTimeOriginal"].strip()
        day = dt[:10].replace(":", "-") if dt else "(no date)"
        year = day[:4] if dt else "(no date)"
        s = sp[sci]
        s["pics"] += 1
        s["files"] += _files(r["Files"])
        s["days"][day] += 1
        s["cams"][cam] += 1
        s["years"][year] += 1
        s["eng"] = s["eng"] or r["EngSoortNaam"].strip()
        s["nl"] = s["nl"] or r["NLSoortNaam"].strip()
        if not s["taxa"][0]:
            s["taxa"] = _taxa(r["TargetPath"])
        if dt:
            s["first"] = min(s["first"] or day, day)
            s["last"] = max(s["last"], day)
        cam_species[cam][sci] += 1
        cam_day[cam][(sci, day)] += 1
        year_cam[year][cam] += 1
        burst[(sci, day, cam)] += 1

    cams = [c for c, _ in Counter({c: sum(v.values()) for c, v in cam_species.items()}).most_common()]
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    def sheet(title, header, data, widths=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
        for row in data:
            ws.append(row)
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, 1):
            w = (widths or {}).get(h) or max(10, min(40, len(str(h)) + 2))
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    # ── Summary ──────────────────────────────────────────────────────────────
    per_species = [s["pics"] for s in sp.values()]
    per_species_day = [n for s in sp.values() for n in s["days"].values()]
    summary = [
        ["Master", str(root)],
        ["Pictures (shots) with a species", sum(per_species)],
        ["Pictures without a species (skipped)", no_species],
        ["Files on disk (RAW+JPG+sidecars)", sum(s["files"] for s in sp.values())],
        ["Species", len(sp)],
        [],
        ["Measure", "n", "total", "mean", "median", "p75", "p90", "p95", "max"],
    ]
    for label, vals in [("Pictures per species", per_species),
                        ("Pictures per species per day", per_species_day)]:
        st = _stats(vals)
        summary.append([label] + [st[k] for k in ("n", "total", "mean", "median", "p75", "p90", "p95", "max")])
    summary += [[], ["Species by number of pictures", "species", "% of species", "pictures", "% of pictures"]]
    tot_p = sum(per_species)
    for lo, hi in BUCKETS:
        hit = [n for n in per_species if lo <= n <= hi]
        label = f"{lo}" if lo == hi else (f"{lo}+" if hi >= 10**9 else f"{lo}–{hi}")
        summary.append([label, len(hit), round(100 * len(hit) / len(sp), 1),
                        sum(hit), round(100 * sum(hit) / tot_p, 1)])
    ws = sheet("Summary", ["What", "Value"], summary, {"What": 38, "Value": 14}, first=True)
    for row in ws.iter_rows():
        if row[0].value in ("Measure", "Species by number of pictures"):
            for c in row:
                c.font, c.fill = bold, head_fill

    # ── Per species ──────────────────────────────────────────────────────────
    hdr = (["Scientific", "English", "Dutch", "Class", "Order", "Family", "Pictures", "Files",
            "Days", "Avg / day", "Median / day", "Max / day", "Best day",
            "First", "Last"] + cams)
    data = []
    for sci, s in sorted(sp.items(), key=lambda kv: -kv[1]["pics"]):
        dvals = list(s["days"].values())
        best_day, best_n = s["days"].most_common(1)[0]
        data.append([sci, s["eng"], s["nl"], *s["taxa"], s["pics"], s["files"], len(dvals),
                     round(statistics.fmean(dvals), 1), statistics.median(dvals), best_n, best_day,
                     s["first"], s["last"]] + [s["cams"].get(c, 0) or None for c in cams])
    sheet("Per species", hdr, data, {"Scientific": 30, "English": 26, "Dutch": 26,
                                     "Class": 14, "Order": 16, "Family": 18, "Best day": 12})

    # ── Per camera ───────────────────────────────────────────────────────────
    cdata = []
    for c in cams:
        a, b = _stats(list(cam_species[c].values())), _stats(list(cam_day[c].values()))
        cdata.append([c, a["total"], a["n"], a["mean"], a["median"], a["p90"], a["max"],
                      b["n"], b["mean"], b["median"], b["p90"], b["max"]])
    sheet("Per camera", ["Camera", "Pictures", "Species", "Avg / species", "Median / species",
                         "p90 / species", "Max / species", "Species-days", "Avg / species-day",
                         "Median / species-day", "p90 / species-day", "Max / species-day"],
          cdata, {"Camera": 22})

    # ── Per year × camera ────────────────────────────────────────────────────
    ydata = []
    for y in sorted(year_cam):
        ydata.append([y, sum(year_cam[y].values())] + [year_cam[y].get(c, 0) or None for c in cams])
    sheet("Per year", ["Year", "Pictures"] + cams, ydata)

    # ── Biggest bursts (one species, one day, one camera) ────────────────────
    bdata = [[sci, sp[sci]["eng"], sp[sci]["nl"], day, cam, n]
             for (sci, day, cam), n in burst.most_common(500)]
    sheet("Biggest days", ["Scientific", "English", "Dutch", "Day", "Camera", "Pictures"],
          bdata, {"Scientific": 30, "English": 26, "Dutch": 26, "Camera": 18})

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    st = _stats(per_species)
    print(f"{len(rows)} rows, {len(sp)} species, {st['total']} pictures; "
          f"mean {st['mean']}, median {st['median']}, p90 {st['p90']}, max {st['max']}")
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
